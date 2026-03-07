import platform
import os
from docx import Document
from datetime import datetime

def save_to_document(text, filename_prefix="aletheon_report"):
    doc= Document()

    doc.add_heading("Aletheon DRUG Output Report", level=1)

    timestamp= datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    doc.add_paragraph(f"Generated on: {timestamp}\n")
    for line in text.split("\n"):
        doc.add_paragraph(line)

    filename=f"{filename_prefix}.docx"
    filepath=os.path.abspath(filename)

    doc.save(filepath)

    os_name=platform.system()
    import subprocess

    if os_name=='Darwin':
        subprocess.run(["open", filepath])
    elif os_name=='Windows':
        os.startfile(filepath)
    else:
        subprocess.run(["xdg-open", filepath])
    return filepath